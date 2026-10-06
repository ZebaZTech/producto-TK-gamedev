from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# EXPORTACIÓN A EXCEL
# ============================================================

def exportar_excel(datos, encabezados, nombre_archivo, carpeta="exports/excel"):
    """
    Exporta información a un archivo Excel.

    datos:
        Lista de filas que serán exportadas.

    encabezados:
        Nombres de las columnas.

    nombre_archivo:
        Nombre del archivo Excel que se generará.
    """

    # Crear la carpeta si no existe.
    ruta_carpeta = Path(carpeta)
    ruta_carpeta.mkdir(parents=True, exist_ok=True)

    ruta_archivo = ruta_carpeta / nombre_archivo

    # Crear libro y hoja.
    libro = Workbook()
    hoja = libro.active
    hoja.title = "GameDev"

    # --------------------------------------------------------
    # Encabezados
    # --------------------------------------------------------

    for columna, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(row=1, column=columna, value=encabezado)

        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center")

    # --------------------------------------------------------
    # Datos
    # --------------------------------------------------------

    for fila, registro in enumerate(datos, start=2):
        for columna, valor in enumerate(registro, start=1):
            hoja.cell(row=fila, column=columna, value=valor)

    # --------------------------------------------------------
    # Configuración de la hoja
    # --------------------------------------------------------

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    # Ajustar automáticamente el ancho de las columnas.
    for columna in hoja.columns:
        maximo = 0
        numero_columna = columna[0].column

        for celda in columna:
            if celda.value is not None:
                longitud = len(str(celda.value))
                maximo = max(maximo, longitud)

        hoja.column_dimensions[
            get_column_letter(numero_columna)
        ].width = min(maximo + 2, 40)

    libro.save(ruta_archivo)

    return str(ruta_archivo)